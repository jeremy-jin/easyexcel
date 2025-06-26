from collections import OrderedDict
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl.workbook import Workbook
from openpyxl.writer.excel import ExcelWriter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .utils import CollectionProxy


class BaseExportHandler(object):
    SHEETS = []

    def __init__(
        self,
        container,
        filename: str,
        data: dict,
        default_titles=False,
    ):
        self.container = container
        self.filename = filename
        self.data = data
        self.wb = Workbook()
        self.default_titles = default_titles

    def register_sheets(self):
        sheets = []
        for i, sheet_cls in enumerate(self.SHEETS):
            sheet_data = self.data.get(sheet_cls.SHEET_NAME, {})
            sheets.append(
                sheet_cls(
                    self.wb,
                    sheet_cls.SHEET_NAME,
                    sheet_data.get("titles"),
                    sheet_data.get("data"),
                    create_sheet=False if i == 0 else True,
                    default_titles=self.default_titles,
                )
            )

        return sheets

    def get_formatted_filename(self):
        return self.filename

    def generate_bytesio_excel(self):
        sheets = self.register_sheets()
        CollectionProxy(sheets).generate_sheet()
        return self.create_bytesio_of_workbook(self.wb)

    @staticmethod
    def convert_workbook_to_bytesio(workbook):
        """Return an in-memory workbook"""

        temp_buffer = BytesIO()
        archive = ZipFile(temp_buffer, "w", ZIP_DEFLATED, allowZip64=True)
        writer = ExcelWriter(workbook, archive)

        try:
            writer.write_data()
        finally:
            archive.close()

        virtual_workbook = temp_buffer.getvalue()
        temp_buffer.close()

        return virtual_workbook

    def create_bytesio_of_workbook(self, workbook):
        return self.convert_workbook_to_bytesio(workbook)


class BaseGenerateExportSheet(object):
    SHEET_FIELDS = OrderedDict()
    TITLE_CELL_RICH_TEXT = {}
    NUMBER_FORMATS = {}

    TITLE_ROW_HEIGHT = 20
    MAIN_DATA_ROW_HEIGHT = 18

    DEFAULT_TITLE_FONT = Font(size=14, bold=True)
    DEFAULT_MAIN_DATA_FONT = Font(size=12)
    DEFAULT_TITLE_FILL = PatternFill("solid", fgColor="dae7ea")
    DEFAULT_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
    DEFAULT_BORDER = Border(
        left=Side(border_style="thin"),
        right=Side(border_style="thin"),
        top=Side(border_style="thin"),
        bottom=Side(border_style="thin"),
    )

    def __init__(
        self,
        wb,
        sheet_name,
        titles,
        data,
        create_sheet=True,
        default_titles=False,
    ):
        self.wb = wb
        if create_sheet:
            self.ws = self.wb.create_sheet(sheet_name)
        else:
            self.ws = self.wb.active
            self.ws.title = sheet_name

        self.titles = (
            titles or [] if not default_titles else list(self.SHEET_FIELDS.values())
        )
        self.data = data

    def format_title_rich_text(self, title):
        pass

    def write_titles(self):
        self.set_row_dimension(1, self.TITLE_ROW_HEIGHT)
        for ind, column in enumerate(self.titles, start=1):
            column_width = float((len(str(column)) + 2) * 2.5)
            self.set_column_dimension(ind, column_width)
            self.write_cell(
                1,
                ind,
                column,
                font=self.DEFAULT_TITLE_FONT,
                fill=self.DEFAULT_TITLE_FILL,
                alignment=self.DEFAULT_ALIGN,
                border=self.DEFAULT_BORDER,
            )

    def write_cell(self, row, col, value, **kwargs):
        cell = self.ws.cell(row=row, column=col, value=value)
        for key, val in kwargs.items():
            setattr(cell, key, val)

    def set_column_dimension(self, col, width):
        column_dimension = self.ws.column_dimensions[get_column_letter(col)]
        column_dimension.width = width

    def set_row_dimension(self, row, height):
        row_dimensions = self.ws.row_dimensions[row]
        row_dimensions.height = height

    def generate_sheet(self):
        self.write_titles()
        self.write_data()

    def write_data(self):
        for row_num, row in enumerate(self.data or [], start=2):
            self.set_row_dimension(row_num, self.MAIN_DATA_ROW_HEIGHT)
            fields = self.SHEET_FIELDS.keys()
            for col_num, field_name in enumerate(fields, start=1):
                value = row.get(field_name)
                value = self.serialize(value)

                number_format = self.NUMBER_FORMATS.get(field_name, None)
                kwargs = {
                    "font": self.DEFAULT_MAIN_DATA_FONT,
                    "alignment": self.DEFAULT_ALIGN,
                    "border": self.DEFAULT_BORDER,
                }
                if number_format:
                    kwargs["number_format"] = number_format

                self.write_cell(
                    row_num,
                    col_num,
                    value,
                    **kwargs,
                )

    @staticmethod
    def serialize(value):
        if isinstance(value, (tuple, list)):
            value = ", ".join(value)

        return value
